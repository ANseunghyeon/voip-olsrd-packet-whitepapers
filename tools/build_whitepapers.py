"""Build the Korean protocol whitepapers from reviewed content and IANA snapshots.

Requires Python 3.10+ and python-docx. PDF export is a separate Word/LibreOffice step.
"""
from pathlib import Path
import json
import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parents[1]
NS = {'i': 'http://www.iana.org/assignments'}

def ref_label(node):
    value=node.get('data','')
    return 'RFC '+value[3:] if node.get('type')=='rfc' and value.startswith('rfc') and value[3:].isdigit() else value

def visible_text(node):
    text=node.text or ''
    for child in node:
        label=visible_text(child)
        if child.tag.rsplit('}',1)[-1]=='xref' and not label: label=ref_label(child)
        text+=label+(child.tail or '')
    return ' '.join(text.split())

def records(file, registry):
    root = ET.parse(ROOT / 'sources' / (file + '.xml')).getroot()
    node = root.find(".//i:registry[@id='%s']" % registry, NS)
    out = []
    for rec in node.findall('i:record', NS):
        row = {}
        for child in rec:
            tag = child.tag.rsplit('}', 1)[-1]
            if tag == 'xref':
                ref = ref_label(child)
                row.setdefault('refs', []).append(ref)
            else:
                row[tag] = visible_text(child)
        out.append(row)
    return out

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd'); el.set(qn('w:fill'), fill); tcPr.append(el)

def font_run(run, size=10, bold=False, mono=False):
    run.font.name = 'Consolas' if mono else 'Malgun Gothic'
    run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '맑은 고딕')

def hyperlink(p, label, url):
    rel = p.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rel)
    r = OxmlElement('w:r'); rp = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '17476A'); rp.append(color)
    fonts = OxmlElement('w:rFonts'); fonts.set(qn('w:ascii'), 'Malgun Gothic'); fonts.set(qn('w:eastAsia'), '맑은 고딕'); rp.append(fonts)
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '17'); rp.append(size)
    r.append(rp); t = OxmlElement('w:t'); t.text = label; r.append(t); h.append(r); p._p.append(h)

def add_table(doc, block):
    heads, rows = block['headers'], block['rows']
    table = doc.add_table(rows=1, cols=len(heads)); table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = block.get('widths', [16.8 / len(heads)] * len(heads))
    for c,w in zip(table.columns, widths): c.width = Cm(w)
    for c,w in zip(table.rows[0].cells, widths): c.width = Cm(w)
    pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        b=OxmlElement('w:'+side); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ('top','left','bottom','right'):
        padding='45' if block.get('dense') and side in ('top','bottom') else '85'
        m=OxmlElement('w:'+side); m.set(qn('w:w'),padding); m.set(qn('w:type'),'dxa'); margins.append(m)
    pr.append(margins)
    for i,vals in enumerate([heads]+rows):
        cells = table.rows[0].cells if i == 0 else table.add_row().cells
        trPr=cells[0]._tc.getparent().get_or_add_trPr()
        no=OxmlElement('w:cantSplit'); trPr.append(no)
        if i == 0:
            repeat=OxmlElement('w:tblHeader'); trPr.append(repeat)
        for j,(cell,value) in enumerate(zip(cells,vals)):
            cell.width=Cm(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(cell, 'E6EDF3' if i==0 else ('F5F7F9' if i%2==0 else 'FFFFFF'))
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0 if block.get('dense') else 1); p.paragraph_format.space_before=Pt(0 if block.get('dense') else 1)
            p.paragraph_format.line_spacing=1.08 if block.get('dense') else 1.12
            if j==0 and len(str(value))<12: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            font_run(p.add_run(str(value)), block.get('size',8.4), i==0)
    gap=doc.add_paragraph(); gap.paragraph_format.space_after=Pt(3); gap.paragraph_format.space_before=Pt(0)
    gap.paragraph_format.line_spacing=0.5; font_run(gap.add_run(''),3)

def build(spec):
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.9); sec.bottom_margin=Cm(1.8); sec.left_margin=Cm(2.1); sec.right_margin=Cm(2.1)
    sec.header_distance=Cm(0.7); sec.footer_distance=Cm(0.8)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3'):
        style=doc.styles[name]; style.font.name='Malgun Gothic'; style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
    normal=doc.styles['Normal']; normal.font.size=Pt(10)
    normal.paragraph_format.line_spacing=1.22; normal.paragraph_format.space_after=Pt(7)
    for name in ['TOC 1','TOC 2']:
        s=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        s.font.name='Malgun Gothic'; s.font.size=Pt(10); s.font.color.rgb=RGBColor(0,0,0)
        s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
        s.paragraph_format.space_after=Pt(5); s.paragraph_format.line_spacing=1.15
    for name,size in [('Title',26),('Subtitle',12),('Heading 1',16),('Heading 2',12),('Heading 3',10.5)]:
        s=doc.styles[name]; s.font.size=Pt(size); s.font.bold=name!='Subtitle'; s.font.italic=False
        s.paragraph_format.space_before=Pt(12 if name!='Title' else 0); s.paragraph_format.space_after=Pt(8)
    p=sec.header.paragraphs[0]; font_run(p.add_run(spec['short_title']),8)
    p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font_run(p.add_run('2026 10 20  |  '),8)
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); p._p.append(field)
    doc.core_properties.title=spec['title']; doc.core_properties.subject=spec['subtitle']
    doc.core_properties.author='Codex'; doc.core_properties.keywords='VoIP, SIP, RTP, RTCP, olsrd, Korean whitepaper'
    blocks=spec['blocks'].copy()
    firstbreak=next(i for i,b in enumerate(blocks) if b['type']=='break')
    blocks[firstbreak:firstbreak]=[{'type':'break'},{'type':'toc'}]
    for b in blocks:
        typ=b['type']
        if typ=='break': doc.add_page_break()
        elif typ=='title': doc.add_paragraph(b['text'],'Title')
        elif typ=='subtitle': doc.add_paragraph(b['text'],'Subtitle')
        elif typ=='h1': doc.add_heading(b['text'],1)
        elif typ=='h2': doc.add_heading(b['text'],2)
        elif typ=='p': doc.add_paragraph(b['text'])
        elif typ=='toc':
            doc.add_heading('목차',2)
            p=doc.add_paragraph()
            begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
            instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=' TOC \\o "1-1" \\h \\z \\u '
            separate=OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'),'separate')
            end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
            for el in [begin,instr,separate,end]:
                rr=OxmlElement('w:r'); rr.append(el); p._p.append(rr)
        elif typ=='table': add_table(doc,b)
        elif typ=='code':
            for line in b['text'].splitlines():
                p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
                font_run(p.add_run(line),8.2,mono=True)
            doc.add_paragraph().paragraph_format.space_after=Pt(2)
        elif typ=='source':
            p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6)
            font_run(p.add_run(b['text']+'\n'),8.5)
            hyperlink(p,'고정 커밋의 소스 원문 열기' if len(b['url'])>100 else b['url'],b['url'])
        else: raise ValueError(typ)
    # The bundled blank DOCX includes title borders. Remove them deterministically.
    for root in (doc.element,doc.styles.element):
        for border in list(root.iter(qn('w:pBdr'))): border.getparent().remove(border)
    path=ROOT/'whitepapers'/(spec['slug']+'.docx'); doc.save(path)
    return path

def expand_registries(spec):
    for key,title,cols,keys,widths in [
        ('sip-parameters-7','부록 A SIP 응답 코드 전체 목록',['코드','IANA 명칭','근거'],['value','description','refs'],[1.4,9.3,6.1]),
        ('rtp-parameters-4','부록 B RTCP 패킷 유형 전체 목록',['값','약어와 의미','근거'],['value','name','refs'],[1.4,9.3,6.1]),
        ('rtp-parameters-8','부록 C RTPFB FMT 등록 목록',['FMT','등록 명칭','근거'],['value','name','refs'],[1.4,9.3,6.1]),
        ('rtp-parameters-9','부록 D PSFB FMT 등록 목록',['FMT','등록 명칭','근거'],['value','name','refs'],[1.4,9.3,6.1]),
    ]:
        file='sip-parameters' if key.startswith('sip') else 'rtp-parameters'
        data=records(file,key); rows=[]
        for r in data:
            def val(k):
                if k=='refs': return ', '.join(r.get(k,[])) or ('RFC 3261' if file=='sip-parameters' else '-')
                if k=='name': return ' '.join(v for v in [r.get('abbrev',''),r.get('name',r.get('description',r.get('long_name','')))] if v)
                return r.get(k,r.get('code',''))
            rows.append([val(k) for k in keys])
        spec['blocks'] += [{'type':'break'},{'type':'h1','text':title},{'type':'p','text':'2026년 10월 20일 내려받은 IANA XML 스냅샷을 그대로 열거한다. 등록됨은 모든 단말의 필수 구현이나 해당 문서의 Standards Track 상태를 뜻하지 않는다. Reserved와 Unassigned는 송신용 메시지 정의가 아니다. ['+('V01' if file=='sip-parameters' else 'V02')+']'}, {'type':'table','headers':cols,'rows':rows,'widths':widths,'dense':True}]
    # All audio payload names, not an arbitrarily selected codec list.
    data=records('rtp-parameters','rtp-parameters-2')
    audio=[r for r in data if r.get('type','').lower()=='audio']
    rows=[[r.get('subtype',r.get('name','')),', '.join(r.get('refs',[])) or '-'] for r in audio]
    spec['blocks'] += [{'type':'break'},{'type':'h1','text':'부록 E 등록된 RTP 오디오 페이로드 전체 목록'}, {'type':'p','text':'각 행은 audio 하위 미디어 유형이다. 각 코덱의 비트스트림 문법과 선택 매개변수까지 하나의 고정 패킷 목록으로 환원할 수는 없다. 실제 PT는 SDP와 해당 페이로드 RFC로 결정한다. 아래 목록의 정의 근거는 함께 저장한 rtp-parameters.xml의 xref에 보존되어 있다. [V02]'}, {'type':'table','headers':['RTP audio 미디어 하위 유형','근거'],'rows':rows,'widths':[7.6,9.2],'dense':True}]
    for file,registry,title,cite in [
        ('stun-parameters','stun-parameters-2','부록 F STUN 메서드 전체 목록','V58'),
        ('rtcp-xr-block-types','rtcp-xr-block-types-1','부록 G RTCP XR 블록 전체 목록','V59'),
        ('rtp-parameters','rtp-parameters-5','부록 H RTCP SDES 항목 전체 목록','V02'),
    ]:
        rs=records(file,registry)
        rows=[[r.get('value',r.get('bt','')),r.get('description',r.get('name','')),', '.join(r.get('refs',[])) or '-'] for r in rs]
        spec['blocks'] += [{'type':'break'},{'type':'h1','text':title},{'type':'p','text':'2026년 10월 20일 IANA 스냅샷의 전체 등록 범위다. Reserved와 Unassigned는 송신할 수 있는 정의된 동작이 아니다. 등록 항목 중 사설·비표준 확장도 있으므로 개별 근거의 지위를 확인한다. ['+cite+']'}, {'type':'table','headers':['값','등록 내용','근거'],'rows':rows,'widths':[2.2,9.2,5.4],'dense':True}]
    asn=json.loads((ROOT/'sources'/'h323_message_inventory.json').read_text(encoding='utf-8'))
    spec['blocks'] += [{'type':'break'},{'type':'h1','text':'부록 I H2250과 H245 메시지 선택지 전체 목록'}, {'type':'p','text':'ITU-T H.225.0 2022년 3월판과 H.245 2022년 3월판의 지정한 최상위 ASN.1 CHOICE를 모두 열거했다. ASN.1 확장 표식은 유형이 아니므로 행으로 세지 않았다. 각 항목의 선택적 필드, 하위 CHOICE, 후속 정오표와 서비스 부속서는 원문을 따른다. H.245의 일부 선택지는 H.223·ATM 등도 지원하며 모두 오디오 VoIP에서 발생하는 것은 아니다. [V53, V54]'}]
    for group in asn['groups']:
        spec['blocks'] += [{'type':'h2','text':group['label']}, {'type':'p','text':group['purpose']+' 원문 인쇄면 '+group['printed_pages']+'쪽.'}, {'type':'table','headers':['ASN.1 메시지 선택지','형식 이름'],'rows':[[x['name'],x['asn_type']] for x in group['choices']],'widths':[8.4,8.4],'size':8.2,'dense':True}]
    # Preserve complete machine-readable inventories in a compact derived JSON.
    inventory={}
    for f,rs in [('sip-parameters',['sip-parameters-2','sip-parameters-6','sip-parameters-7']),('rtp-parameters',['rtp-parameters-1','rtp-parameters-2','rtp-parameters-4','rtp-parameters-5','rtp-parameters-8','rtp-parameters-9']),('stun-parameters',['stun-parameters-2','stun-parameters-4','stun-parameters-6']),('rtcp-xr-block-types',['rtcp-xr-block-types-1'])]:
        for registry in rs: inventory[registry]=records(f,registry)
    (ROOT/'sources'/'packet_inventory.json').write_text(json.dumps(inventory,ensure_ascii=False,indent=2),encoding='utf-8')
    return spec

if __name__=='__main__':
    for name in ['voip','olsrd']:
        spec=json.loads((ROOT/'content'/(name+'.json')).read_text(encoding='utf-8'))
        if name=='voip':
            refs=spec['blocks'].pop()
            assert refs['type']=='references'
            spec=expand_registries(spec)
            spec['blocks'] += [{'type':'break'},{'type':'h1','text':'참고 문헌'}]+refs['items']
        print(build(spec))
